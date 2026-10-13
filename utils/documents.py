from io import BytesIO
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_DOCUMENT_CHARS = 200_000
MAX_INSTRUCTION_CHARS = 20_000
TEXT_EXTENSIONS = {'.txt', '.md', '.csv', '.json'}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {'.pdf', '.docx'}


def extract_document_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ', '.join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f'Unsupported file type. Use one of: {supported}.')

    try:
        if extension in TEXT_EXTENSIONS:
            text = content.decode('utf-8-sig')
        elif extension == '.pdf':
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted and not reader.decrypt(''):
                raise ValueError('Password-protected PDFs are not supported.')
            chunks = []
            total_chars = 0
            for page in reader.pages:
                page_text = page.extract_text() or ''
                chunks.append(page_text)
                total_chars += len(page_text)
                if total_chars >= MAX_DOCUMENT_CHARS:
                    break
            text = '\n'.join(chunks)
        else:
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(
                ' '.join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            text = '\n'.join(parts)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('Could not read this document. Check that the file is valid and not corrupted.') from error

    text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\x0c', '\n').replace('\xa0', ' ')
    text = re.sub(r'[\x00-\x08\x0b\x0e-\x1f\x7f]', '', text)
    text = re.sub(r'[ \t]+\n', '\n', text)
    text = re.sub(r'\n[ \t]*\n(?:[ \t]*\n)+', '\n\n', text)
    text = text.strip()
    if not text:
        raise ValueError('No readable text was found in this document.')
    if len(text) > MAX_DOCUMENT_CHARS:
        text = text[:MAX_DOCUMENT_CHARS] + f'\n\n[Document text truncated at {MAX_DOCUMENT_CHARS:,} characters.]'
    return text